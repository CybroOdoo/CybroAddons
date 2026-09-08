# -*- coding: utf-8 -*-
#############################################################################
#
#    Cybrosys Technologies Pvt. Ltd.
#
#    Copyright (C) 2026-TODAY Cybrosys Technologies(<https://www.cybrosys.com>)
#    Author: Cybrosys Techno Solutions (odoo@cybrosys.com)
#
#    You can modify it under the terms of the GNU LESSER
#    GENERAL PUBLIC LICENSE (LGPL v3), Version 3.
#
#    This program is distributed in the hope that it will be useful,
#    but WITHOUT ANY WARRANTY; without even the implied warranty of
#    MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
#    GNU LESSER GENERAL PUBLIC LICENSE (LGPL v3) for more details.
#
#    You should have received a copy of the GNU LESSER GENERAL PUBLIC LICENSE
#    (LGPL v3) along with this program.
#    If not, see <https://www.gnu.org/licenses/>.
#
#############################################################################
import base64
import io

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from odoo import api, fields, models, _
from odoo.exceptions import AccessError, UserError, ValidationError


class ComplianceReport(models.TransientModel):
    """Customizable regulatory compliance report builder supporting multi-criteria filtering."""
    _name = 'wm.compliance.report'
    _description = 'Compliance Report Builder'

    report_type = fields.Selection([
        ('manifest', 'Waste Manifests'),
        ('certificate', 'Compliance Certificates'),
    ], string='Report Type', required=True, default='manifest')

    date_from = fields.Date(string='Date From', required=True, default=fields.Date.context_today)
    date_to = fields.Date(string='Date To', required=True, default=fields.Date.context_today)
    facility_id = fields.Many2one('wm.disposal.facility', string='Disposal Facility')
    output_format = fields.Selection([
        ('pdf', 'PDF Report'),
        ('xlsx', 'Excel spreadsheet (XLSX)')
    ], string='Output Format', required=True, default='pdf')
    company_id = fields.Many2one(
        'res.company',
        string='Company',
        required=True,
        default=lambda self: self.env.company
    )

    @api.model
    def default_get(self, fields_list):
        """
        Verify that the user possesses read access to disposal facilities
        prior to initializing the report configuration wizard.
        """
        res = super().default_get(fields_list)
        if not self.env['wm.disposal.facility'].has_access('read'):
            raise AccessError(_("You do not have access rights to the Disposal Facility and cannot access this Compliance Report."))
        return res

    @api.model_create_multi
    def create(self, vals_list):
        """
        Enforce disposal facility read authorization before allowing record creation.
        """
        if not self.env['wm.disposal.facility'].has_access('read'):
            raise AccessError(_("You do not have access rights to the Disposal Facility and cannot access this Compliance Report."))
        return super().create(vals_list)

    @api.constrains('date_from', 'date_to')
    def _check_date_range(self):
        """
        Validate that Date To is not earlier than Date From.
        """
        for record in self:
            if record.date_from and record.date_to and record.date_to < record.date_from:
                raise ValidationError(_("End Date cannot be earlier than Start Date."))

    def action_generate(self):
        """
        Compile the configured compliance report dataset from the selected
        records and render the output as a downloadable PDF or Excel file attachment.
        """
        self.ensure_one()
        if not self.env['wm.disposal.facility'].has_access('read'):
            raise AccessError(_("You do not have access rights to the Disposal Facility and cannot access this Compliance Report."))
        if self.date_to < self.date_from:
            raise UserError(_("End Date cannot be earlier than Start Date."))

        if self.report_type == 'certificate':
            return self._generate_certificate_report()
        else:
            return self._generate_manifest_report()

    def _generate_manifest_report(self):
        """Generate PDF or XLSX report for waste compliance manifests."""
        domain = [
            ('manifest_date', '>=', self.date_from),
            ('manifest_date', '<=', self.date_to),
            ('company_id', '=', self.company_id.id),
        ]
        if self.facility_id:
            domain.append(('disposal_facility_id', '=', self.facility_id.id))

        manifests = self.env['wm.compliance.manifest'].search(domain)
        if not manifests:
            raise UserError(_(
                "No Manifests found for the specified criteria. "
                "Please adjust the date range or facility filter."
            ))

        if self.output_format == 'pdf':
            report = self.env.ref('wm_compliance.action_report_compliance_manifest')
            return report.report_action(manifests.ids, config=False)
        else:
            return self._generate_manifest_xlsx(manifests)

    def _generate_certificate_report(self):
        """Generate PDF or XLSX report for compliance certificates."""
        start_dt = fields.Datetime.to_datetime(self.date_from)
        end_dt = fields.Datetime.to_datetime(self.date_to).replace(hour=23, minute=59, second=59)

        domain = [
            ('company_id', '=', self.company_id.id),
            '|',
            '&', ('issued_date', '>=', self.date_from), ('issued_date', '<=', self.date_to),
            '&', ('create_date', '>=', start_dt), ('create_date', '<=', end_dt),
        ]
        if self.facility_id:
            domain.append(('facility_id', '=', self.facility_id.id))

        certificates = self.env['wm.compliance.certificate'].search(domain)
        if not certificates:
            raise UserError(_(
                "No Compliance Certificates found for the specified criteria. "
                "Please adjust the date range or facility filter."
            ))

        if self.output_format == 'pdf':
            report = self.env.ref('wm_compliance.action_report_compliance_certificate')
            return report.report_action(certificates.ids, config=False)
        else:
            return self._generate_certificate_xlsx(certificates)

    def _generate_manifest_xlsx(self, manifests):
        """Build styled Excel workbook for manifests."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Manifests Report"

        GREEN_DARK = "1B4D3E"
        GREEN_MID = "2E7D5A"
        GREEN_LIGHT = "D6ECD9"
        WHITE = "FFFFFF"
        GREY_TEXT = "555555"
        AMBER = "FF8C00"

        thin = Side(style='thin', color="AAAAAA")
        thin_border = Border(left=thin, right=thin, top=thin, bottom=thin)

        COL_HEADERS = [
            "Manifest Ref", "Generator", "Transporter",
            "Disposal Facility", "Total Qty", "State", "Date",
        ]
        COL_WIDTHS = [20, 28, 28, 28, 13, 14, 14]
        NUM_COLS = len(COL_HEADERS)

        # Title row
        title_text = (
            f"Compliance Manifests Report  |  "
            f"{self.date_from}  →  {self.date_to}"
            + (f"  |  {self.facility_id.name}" if self.facility_id else "")
        )
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=NUM_COLS)
        title_cell = ws.cell(row=1, column=1, value=title_text)
        title_cell.font = Font(name="Calibri", bold=True, size=14, color=WHITE)
        title_cell.fill = PatternFill("solid", fgColor=GREEN_DARK)
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28

        # Meta row
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=NUM_COLS)
        meta_cell = ws.cell(
            row=2, column=1,
            value=f"Company: {self.company_id.name or ''}    "
                  f"Generated: {fields.Date.context_today(self)}"
        )
        meta_cell.font = Font(name="Calibri", italic=True, size=10, color=WHITE)
        meta_cell.fill = PatternFill("solid", fgColor=GREEN_MID)
        meta_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 18

        ws.row_dimensions[3].height = 6

        # Column headers
        header_font = Font(name="Calibri", bold=True, size=11, color=WHITE)
        header_fill = PatternFill("solid", fgColor=GREEN_MID)
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col_idx, header in enumerate(COL_HEADERS, start=1):
            cell = ws.cell(row=4, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border
        ws.row_dimensions[4].height = 22

        data_font_base = Font(name="Calibri", size=10)
        fill_even = PatternFill("solid", fgColor=GREEN_LIGHT)
        fill_odd = PatternFill("solid", fgColor=WHITE)
        centre = Alignment(horizontal="center", vertical="center")
        left = Alignment(horizontal="left", vertical="center", wrap_text=True)

        STATE_COLOURS = {
            'draft': "999999",
            'submitted': "3A7BD5",
            'collected': "F59E0B",
            'disposed': "2E7D5A",
            'cancelled': "C0392B",
        }

        total_qty = 0.0
        for row_offset, m in enumerate(manifests):
            excel_row = 5 + row_offset
            fill = fill_even if row_offset % 2 == 0 else fill_odd

            row_data = [
                m.name or '',
                m.generator_id.name or '',
                m.transporter_id.name or '',
                m.disposal_facility_id.name or '',
                m.total_quantity or 0,
                (m.state or '').replace('_', ' ').title(),
                str(m.manifest_date) if m.manifest_date else '',
            ]
            total_qty += m.total_quantity or 0

            for col_idx, value in enumerate(row_data, start=1):
                cell = ws.cell(row=excel_row, column=col_idx, value=value)
                cell.font = data_font_base
                cell.fill = fill
                cell.border = thin_border
                if col_idx in (5,):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = '#,##0.00'
                elif col_idx in (6, 7):
                    cell.alignment = centre
                    if col_idx == 6:
                        state_key = (m.state or '')
                        colour = STATE_COLOURS.get(state_key, GREY_TEXT)
                        cell.font = Font(name="Calibri", size=10, bold=True, color=colour)
                else:
                    cell.alignment = left
            ws.row_dimensions[excel_row].height = 16

        # Totals row
        total_row = 5 + len(manifests)
        totals_fill = PatternFill("solid", fgColor="FFF3CD")
        totals_font = Font(name="Calibri", bold=True, size=10, color=AMBER)
        label_cell = ws.cell(row=total_row, column=1, value="TOTAL")
        label_cell.font = totals_font
        label_cell.fill = totals_fill
        label_cell.border = thin_border
        label_cell.alignment = Alignment(horizontal="right")
        for col_idx in range(2, NUM_COLS + 1):
            cell = ws.cell(row=total_row, column=col_idx)
            cell.fill = totals_fill
            cell.border = thin_border
            if col_idx == 5:
                cell.value = total_qty
                cell.font = totals_font
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '#,##0.00'
        ws.row_dimensions[total_row].height = 18

        for col_idx, width in enumerate(COL_WIDTHS, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        ws.freeze_panes = "A5"
        ws.auto_filter.ref = f"A4:{get_column_letter(NUM_COLS)}{total_row - 1}"

        fp = io.BytesIO()
        wb.save(fp)
        excel_bytes = fp.getvalue()
        fp.close()

        filename = f"Compliance_Manifests_{self.date_from}_to_{self.date_to}.xlsx"
        attachment = self.env['ir.attachment'].create({
            'name': filename,
            'type': 'binary',
            'datas': base64.b64encode(excel_bytes),
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            'res_model': 'wm.compliance.report',
            'res_id': self.id,
        })

        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'new',
        }

    def _generate_certificate_xlsx(self, certificates):
        """Build styled Excel workbook for compliance certificates."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Certificates Report"

        GREEN_DARK = "1B4D3E"
        GREEN_MID = "2E7D5A"
        GREEN_LIGHT = "D6ECD9"
        WHITE = "FFFFFF"
        GREY_TEXT = "555555"
        AMBER = "FF8C00"

        thin = Side(style='thin', color="AAAAAA")
        thin_border = Border(left=thin, right=thin, top=thin, bottom=thin)

        COL_HEADERS = [
            "Certificate No", "Certificate Name", "Operation Type",
            "Disposal Facility", "Quantity", "UoM", "State", "Issued Date", "Signature Hash"
        ]
        COL_WIDTHS = [20, 26, 18, 28, 14, 10, 14, 14, 32]
        NUM_COLS = len(COL_HEADERS)

        # Title row
        title_text = (
            f"Compliance Certificates Report  |  "
            f"{self.date_from}  →  {self.date_to}"
            + (f"  |  {self.facility_id.name}" if self.facility_id else "")
        )
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=NUM_COLS)
        title_cell = ws.cell(row=1, column=1, value=title_text)
        title_cell.font = Font(name="Calibri", bold=True, size=14, color=WHITE)
        title_cell.fill = PatternFill("solid", fgColor=GREEN_DARK)
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28

        # Meta row
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=NUM_COLS)
        meta_cell = ws.cell(
            row=2, column=1,
            value=f"Company: {self.company_id.name or ''}    "
                  f"Generated: {fields.Date.context_today(self)}"
        )
        meta_cell.font = Font(name="Calibri", italic=True, size=10, color=WHITE)
        meta_cell.fill = PatternFill("solid", fgColor=GREEN_MID)
        meta_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 18

        ws.row_dimensions[3].height = 6

        # Column headers
        header_font = Font(name="Calibri", bold=True, size=11, color=WHITE)
        header_fill = PatternFill("solid", fgColor=GREEN_MID)
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col_idx, header in enumerate(COL_HEADERS, start=1):
            cell = ws.cell(row=4, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border
        ws.row_dimensions[4].height = 22

        data_font_base = Font(name="Calibri", size=10)
        fill_even = PatternFill("solid", fgColor=GREEN_LIGHT)
        fill_odd = PatternFill("solid", fgColor=WHITE)
        centre = Alignment(horizontal="center", vertical="center")
        left = Alignment(horizontal="left", vertical="center", wrap_text=True)

        STATE_COLOURS = {
            'draft': "999999",
            'issued': "2E7D5A",
        }

        total_qty = 0.0
        for row_offset, cert in enumerate(certificates):
            excel_row = 5 + row_offset
            fill = fill_even if row_offset % 2 == 0 else fill_odd

            row_data = [
                cert.certificate_no or '',
                cert.name or '',
                (cert.certificate_type or '').capitalize(),
                cert.facility_id.name or '',
                cert.quantity or 0.0,
                cert.uom_id.name or '',
                (cert.state or '').replace('_', ' ').title(),
                str(cert.issued_date) if cert.issued_date else '',
                cert.signature_pdf_hash or ('Signed' if cert.signature_id else 'Unsigned'),
            ]
            total_qty += cert.quantity or 0.0

            for col_idx, value in enumerate(row_data, start=1):
                cell = ws.cell(row=excel_row, column=col_idx, value=value)
                cell.font = data_font_base
                cell.fill = fill
                cell.border = thin_border
                if col_idx in (5,):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = '#,##0.00'
                elif col_idx in (1, 3, 6, 7, 8):
                    cell.alignment = centre
                    if col_idx == 7:
                        state_key = (cert.state or '')
                        colour = STATE_COLOURS.get(state_key, GREY_TEXT)
                        cell.font = Font(name="Calibri", size=10, bold=True, color=colour)
                else:
                    cell.alignment = left
            ws.row_dimensions[excel_row].height = 16

        # Totals row
        total_row = 5 + len(certificates)
        totals_fill = PatternFill("solid", fgColor="FFF3CD")
        totals_font = Font(name="Calibri", bold=True, size=10, color=AMBER)
        label_cell = ws.cell(row=total_row, column=1, value="TOTAL")
        label_cell.font = totals_font
        label_cell.fill = totals_fill
        label_cell.border = thin_border
        label_cell.alignment = Alignment(horizontal="right")
        for col_idx in range(2, NUM_COLS + 1):
            cell = ws.cell(row=total_row, column=col_idx)
            cell.fill = totals_fill
            cell.border = thin_border
            if col_idx == 5:
                cell.value = total_qty
                cell.font = totals_font
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '#,##0.00'
        ws.row_dimensions[total_row].height = 18

        for col_idx, width in enumerate(COL_WIDTHS, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        ws.freeze_panes = "A5"
        ws.auto_filter.ref = f"A4:{get_column_letter(NUM_COLS)}{total_row - 1}"

        fp = io.BytesIO()
        wb.save(fp)
        excel_bytes = fp.getvalue()
        fp.close()

        filename = f"Compliance_Certificates_{self.date_from}_to_{self.date_to}.xlsx"
        attachment = self.env['ir.attachment'].create({
            'name': filename,
            'type': 'binary',
            'datas': base64.b64encode(excel_bytes),
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            'res_model': 'wm.compliance.report',
            'res_id': self.id,
        })

        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'new',
        }
