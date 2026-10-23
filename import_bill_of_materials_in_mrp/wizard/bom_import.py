# -*- coding: utf-8 -*-
#############################################################################
#
#    Cybrosys Technologies Pvt. Ltd.
#
#    Copyright (C) 2026-TODAY Cybrosys Technologies(<https://www.cybrosys.com>)
#    Author: Cybrosys Techno Solutions(<https://www.cybrosys.com>)
#
#    You can modify it under the terms of the GNU LESSER
#    GENERAL PUBLIC LICENSE (LGPL v3), Version 3.
#
#    This program is distributed in the hope that it will be useful,
#    but WITHOUT ANY WARRANTY; without even the implied warranty of
#    MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
#    GNU LESSER GENERAL PUBLIC LICENSE (LGPL v3) for more details.
#
#############################################################################
import base64
import csv
from io import BytesIO
import logging
import zipfile
import zlib

import openpyxl
import xlrd

from odoo import fields, models
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)


def is_binary_file(file_bytes):
    """Check if byte sequence starts with zip/excel header."""
    if file_bytes.startswith(b'PK\x03\x04') or file_bytes.startswith(b'\xd0\xcf\x11\xe0'):
        return True
    return False


def _extract_file_content(file_bytes):
    """If file_bytes is a zip archive, extract inner spreadsheet/csv file."""
    if zipfile.is_zipfile(BytesIO(file_bytes)):
        with zipfile.ZipFile(BytesIO(file_bytes)) as z:
            namelist = z.namelist()
            if 'xl/workbook.xml' in namelist or '[Content_Types].xml' in namelist:
                return file_bytes
            for name in namelist:
                if (
                    not name.startswith('__MACOSX')
                    and not name.startswith('.')
                    and not z.getinfo(name).is_dir()
                ):
                    if name.lower().endswith(('.xlsx', '.xls', '.csv')):
                        return z.read(name)
            for name in namelist:
                if (
                    not name.startswith('__MACOSX')
                    and not name.startswith('.')
                    and not z.getinfo(name).is_dir()
                ):
                    return z.read(name)
    return file_bytes


def _read_excel_rows(file_bytes):
    """Read rows from Excel file (.xlsx or .xls). Returns list of lists/tuples of values."""
    errors = []

    # 1. Try openpyxl read_only=True (.xlsx)
    try:
        wb = openpyxl.load_workbook(filename=BytesIO(file_bytes), data_only=True, read_only=True)
        ws = wb.active
        rows = [list(row) for row in ws.iter_rows(values_only=True)]
        if rows:
            return rows
    except Exception as e:
        errors.append("openpyxl (read_only=True): %s" % e)

    # 2. Try openpyxl read_only=False (.xlsx)
    try:
        wb = openpyxl.load_workbook(filename=BytesIO(file_bytes), data_only=False, read_only=False)
        ws = wb.active
        rows = [[cell.value for cell in row] for row in ws.rows]
        if rows:
            return rows
    except Exception as e:
        errors.append("openpyxl (read_only=False): %s" % e)

    # 3. Try xlrd (.xls)
    try:
        wb = xlrd.open_workbook(file_contents=file_bytes)
        ws = wb.sheet_by_index(0)
        rows = [ws.row_values(r) for r in range(ws.nrows)]
        if rows:
            return rows
    except Exception as e:
        errors.append("xlrd: %s" % e)

    raise ValueError(f"Could not parse Excel file. Errors: {'; '.join(errors)}")


def _read_csv_rows(file_bytes):
    """Read rows from CSV file supporting UTF-8, UTF-8-SIG, UTF-16 (with BOM), Latin1."""
    if is_binary_file(file_bytes):
        raise ValueError("Uploaded file is a binary file (Excel/ZIP), not a valid text CSV file.")

    content = None

    # Check for explicit UTF-16 Byte Order Mark (BOM)
    if file_bytes.startswith((b'\xff\xfe', b'\xfe\xff')):
        for encoding in ('utf-16', 'utf-16-le', 'utf-16-be'):
            try:
                decoded = file_bytes.decode(encoding)
                if '\x00' not in decoded:
                    content = decoded
                    break
            except Exception:
                continue

    # Try standard UTF-8-SIG and UTF-8 first for standard CSV files
    if content is None:
        for encoding in ('utf-8-sig', 'utf-8'):
            try:
                content = file_bytes.decode(encoding)
                break
            except UnicodeDecodeError:
                continue

    if content is None:
        content = file_bytes.decode('latin1', errors='replace').replace('\x00', '')

    content = content.lstrip('\ufeff\ufffe')
    reader = csv.reader(content.splitlines())
    rows = []
    for row in reader:
        cleaned_row = [str(cell).strip('\ufeff\ufffe\x00 ') for cell in row]
        if any(cleaned_row):
            rows.append(cleaned_row)
    return rows


def _get_val(col, idx, default=''):
    if idx < len(col) and col[idx] is not None:
        val = col[idx]
        if isinstance(val, float) and val.is_integer():
            val = int(val)
        val_str = str(val).strip('\ufeff\ufffe\x00 ')
        if val_str:
            return val_str
    return default


def safe_get_file_bytes(data):
    """Safely extract raw file bytes from Odoo 20 BinaryValueAttachment,
    bytes, base64 strings, or zlib-compressed attachments.
    """
    if not data:
        return b''

    raw_bytes = None
    if isinstance(data, str):
        data_str = data.strip()
        if ',' in data_str and 'data:' in data_str[:30]:
            data_str = data_str.split(',', 1)[1]
        missing_padding = len(data_str) % 4
        if missing_padding:
            data_str += '=' * (4 - missing_padding)
        try:
            raw_bytes = base64.b64decode(data_str)
        except Exception:
            raw_bytes = data_str.encode('utf-8', errors='ignore')
    else:
        try:
            raw_bytes = bytes(data)
        except Exception:
            raw_bytes = str(data).encode('utf-8', errors='ignore')

    if not raw_bytes:
        return b''

    # Check if raw_bytes is zlib-compressed data (starts with 0x78)
    if raw_bytes.startswith((b'\x78\x01', b'\x78\x9c', b'\x78\xda')):
        try:
            return zlib.decompress(raw_bytes)
        except Exception as e:
            _logger.debug("Zlib raw decompression skipped: %s", e)

    # Secondary check: if raw_bytes was a base64 string that needs decoding
    if not (raw_bytes.startswith(b'PK\x03\x04') or raw_bytes.startswith(b'\xd0\xcf\x11\xe0')):
        try:
            data_str = raw_bytes.decode('ascii').strip()
            if ',' in data_str and 'data:' in data_str[:30]:
                data_str = data_str.split(',', 1)[1]
            missing_padding = len(data_str) % 4
            if missing_padding:
                data_str += '=' * (4 - missing_padding)
            decoded = base64.b64decode(data_str, validate=True)
            if decoded.startswith((b'\x78\x01', b'\x78\x9c', b'\x78\xda')):
                try:
                    return zlib.decompress(decoded)
                except Exception as e:
                    _logger.debug("Zlib decoded decompression skipped: %s", e)
            return decoded
        except Exception as e:
            _logger.debug("Secondary base64 decode check skipped: %s", e)

    return raw_bytes


class BomImport(models.TransientModel):
    """This is a transient model for creating a wizard for importing bom """
    _name = 'bom.import'
    _description = 'Import Bom wizard'

    file_type = fields.Selection([('csv', 'CSV File'),
                                  ('xls', 'Excel File')],
                                 required=True,
                                 default='csv',
                                 help='File type of importing file')
    name = fields.Char()
    bom_type = fields.Selection(
        [('mtp', 'Manufacture This Product'), ('kit', 'Kit')],
        string='BOM Type', required=True,
        default='mtp', help='Type of Bill of Materials')
    product_variant_by = fields.Selection(
        [('default_code', 'Internal Reference'),
         ('barcode', 'Barcode')], required=True,
        default='default_code', help='Product variant type in file')
    product_by = fields.Selection([('name', 'Name'),
                                   ('default_code', 'Internal Reference'),
                                   ('barcode', 'Barcode')],
                                  required=True, default='name',
                                  help='Product type in file')
    file = fields.Binary(string='Upload File', required=True,
                         help='File to Import')

    def action_import_bom(self):
        """function for importing bom through csv or Excel file"""
        for rec in self:
            if not rec.file:
                continue

            file_bytes = safe_get_file_bytes(rec.file)
            extracted_bytes = _extract_file_content(file_bytes)

            rows = []
            if rec.file_type == 'xls':
                msg = (
                    "Sorry, The Excel file you provided does not match our required format\n"
                    "Error: %s"
                )
                try:
                    rows = _read_excel_rows(extracted_bytes)
                except Exception as e:
                    if not is_binary_file(extracted_bytes):
                        try:
                            rows = _read_csv_rows(extracted_bytes)
                        except Exception:
                            raise UserError(rec.env._(msg, str(e)))
                    else:
                        raise UserError(rec.env._(msg, str(e)))
            else:
                msg = (
                    "Sorry, The CSV file you provided does not match our required format\n"
                    "Error: %s"
                )
                try:
                    rows = _read_csv_rows(extracted_bytes)
                except Exception as e:
                    if is_binary_file(extracted_bytes):
                        try:
                            rows = _read_excel_rows(extracted_bytes)
                        except Exception:
                            raise UserError(rec.env._(msg, str(e)))
                    else:
                        raise UserError(rec.env._(msg, str(e)))

            valid_rows = []
            for r in rows:
                if any(r):
                    has_val = any(c is not None and str(c).strip('\ufeff\ufffe\x00 ') for c in r)
                    if has_val:
                        valid_rows.append(r)

            if not valid_rows or len(valid_rows) < 2:
                raise UserError(rec.env._("The file contains no data to import!"))

            rec_count = 0
            last_bom_id = False

            for col in valid_rows[1:]:
                ref = _get_val(col, 0)
                finished_prod = _get_val(col, 1)
                finished_variant = _get_val(col, 2)
                qty = _get_val(col, 3, default='1')
                uom_name = _get_val(col, 4)
                material_variant = _get_val(col, 5)
                material_qty = _get_val(col, 6, default='1')
                material_uom_name = _get_val(col, 7)

                product = False
                variant = False

                if finished_prod:
                    rec_count += 1
                    product_search = rec.product_by
                    product = self.env['product.template'].search(
                        [(product_search, '=', finished_prod)], limit=1
                    )
                    if not product:
                        product = self.env['product.template'].search([
                            '|', '|',
                            ('default_code', '=', finished_prod),
                            ('name', '=', finished_prod),
                            ('barcode', '=', finished_prod)
                        ], limit=1)

                    if not product:
                        raise UserError(
                            rec.env._("Product template '%s' not found!", finished_prod)
                        )

                    bom_data = {
                        'product_tmpl_id': product.id,
                        'code': ref,
                        'type': 'normal' if rec.bom_type == 'mtp' else 'phantom',
                    }

                    try:
                        bom_data['product_qty'] = float(qty)
                    except ValueError:
                        bom_data['product_qty'] = 1.0

                    if finished_variant:
                        variant_search = rec.product_variant_by
                        variant = self.env['product.product'].search(
                            [(variant_search, '=', finished_variant)], limit=1
                        )
                        if not variant:
                            variant = self.env['product.product'].search([
                                '|', '|',
                                ('default_code', '=', finished_variant),
                                ('barcode', '=', finished_variant),
                                ('name', '=', finished_variant)
                            ], limit=1)
                        if variant:
                            bom_data['product_id'] = variant.id

                    if uom_name:
                        uom = self.env['uom.uom'].search([('name', '=', uom_name)], limit=1)
                        if uom:
                            bom_data['uom_id'] = uom.id

                    if 'uom_id' not in bom_data:
                        if variant and variant.uom_id:
                            bom_data['uom_id'] = variant.uom_id.id
                        elif product and product.uom_id:
                            bom_data['uom_id'] = product.uom_id.id

                    bom_bom = self.env['mrp.bom'].create(bom_data)
                    last_bom_id = bom_bom.id

                if material_variant and last_bom_id:
                    variant_search = rec.product_variant_by
                    material = self.env['product.product'].search(
                        [(variant_search, '=', material_variant)], limit=1
                    )
                    if not material:
                        material = self.env['product.product'].search([
                            '|', '|',
                            ('default_code', '=', material_variant),
                            ('barcode', '=', material_variant),
                            ('name', '=', material_variant)
                        ], limit=1)

                    if not material:
                        raise UserError(
                            rec.env._("Material product variant '%s' not found!", material_variant)
                        )

                    bom_line_dict = {
                        'product_id': material.id,
                        'bom_id': last_bom_id,
                    }

                    try:
                        bom_line_dict['product_qty'] = float(material_qty)
                    except ValueError:
                        bom_line_dict['product_qty'] = 1.0

                    if material_uom_name:
                        uom = self.env['uom.uom'].search(
                            [('name', '=', material_uom_name)], limit=1
                        )
                        if uom:
                            bom_line_dict['uom_id'] = uom.id

                    if 'uom_id' not in bom_line_dict:
                        if material and material.uom_id:
                            bom_line_dict['uom_id'] = material.uom_id.id

                    self.env['mrp.bom.line'].create(bom_line_dict)

            return rec.success_message(rec_count)

    def success_message(self, rec_count):
        """function for displaying success message"""
        msg = f"{rec_count} Records imported successfully"
        message_id = self.env['success.message'].create({'message': msg})
        return {
            'name': 'Message',
            'type': 'ir.actions.act_window',
            'view_mode': 'form',
            'res_model': 'success.message',
            'res_id': message_id.id,
            'target': 'new'
        }
